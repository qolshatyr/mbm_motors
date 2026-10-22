#!/usr/bin/env python3
"""Classify a CloudShop product export by its names: category, part brand, car fitment, years.

    python3 tools/catalog/classify_products.py docs/catalog-<id>.xlsx docs/catalog/Catalog_Classification.xlsx

The export has no filled category, brand or country columns, so everything is derived from the
product name, e.g. «Г3-4 Радиатор кондиционера MBM-AG HY ACCENT/SOLARIS RIO 11-». Results are
proposals for review: every row says what was recognised and what needs a human check. Purchase
prices and costs from the export are not copied.
"""
import re
import sys
from collections import Counter

if __package__:
    from .catalog_rules import (BRANDS, COMPOUNDS, COUNTRIES, ENGINE_FAMILIES, GENERATIONS, GROUPS, MAKE_TOKENS,
                                MODELS, NOISE, PART_RULES)
else:  # run as a script
    sys.path.insert(0, __import__('os').path.dirname(__file__))
    from catalog_rules import (BRANDS, COMPOUNDS, COUNTRIES, ENGINE_FAMILIES, GENERATIONS, GROUPS, MAKE_TOKENS,
                               MODELS, NOISE, PART_RULES)

SHELF = re.compile(r'^\s*([A-Za-zА-Яа-яЁё]{1,2})\s?-?\s?(\d{1,2})\s?-\s?(\d{1,2})(?=\s|$)')
BOX = re.compile(r'\bЯЩИК\s?-?\s?(\d+)\b', re.IGNORECASE)
YEAR_RANGE = re.compile(r'^\(?((?:19|20)?\d{2})-((?:19|20)?\d{2})?\)?$')
TO_YEAR = re.compile(r'^\(?-((?:19|20)\d{2}|\d{2})\)?$')
ENGINE_CODE = re.compile(r'^[GDL]\d[A-Z]{2}$')
VOLUME = re.compile(r'^[1-3]\.\d$')
TOKEN = re.compile(r"[A-Za-z0-9][A-Za-z0-9'.\-]*|-\d{2,4}\)?|[А-Яа-яЁё]+|[()/,+*]")
DISTINCT_GENERATIONS = {code for code in GENERATIONS if len(code) >= 3}
MAX_YEAR = 2027


CYR_TO_LAT = str.maketrans('АВЕКМНОРСТХаеорсухкм', 'ABEKMHOPCTXaeopcyxkm')
LAT_TO_CYR = str.maketrans('ABEKMHOPCTXaeopcyxkm', 'АВЕКМНОРСТХаеорсухкм')


def unmix(word):
    """Fix words typed with look-alike letters of the other script: «Fе», «мacлянoгo»."""
    latin = len(re.findall(r'[A-Za-z]', word))
    cyrillic = len(re.findall(r'[А-Яа-яЁё]', word))
    if not latin or not cyrillic:
        return word
    return word.translate(CYR_TO_LAT) if latin >= cyrillic else word.translate(LAT_TO_CYR)


def normalize(name):
    text = re.sub(r'\s+', ' ', str(name)).strip()
    text = text.replace('→', '-').replace('–', '-').replace('—', '-').replace('->', '-').replace('●', ' ')
    # Split words glued across scripts first: «передHYUNDAI», «салонаSANTA», «ГидронасосOSSCA».
    text = re.sub(r'(?<=[а-яё])(?=[A-Z]{2})|(?<=[A-Z]{2})(?=[А-Яа-яЁё]{2})', ' ', text)
    text = ' '.join(unmix(word) for word in text.split(' '))
    text = re.sub(r'(?<=\d)>', '-', text)
    text = re.sub(r'\b(\d{2}-)(\d\.\d)\b', r'\1 \2', text)
    # «LKIA», «RMOBI»: a side letter glued to a make or brand.
    text = re.sub(r'\b([LR])(KIA|HY|HYUNDAI|MOBI|MOBIS|HPT|LEON)\b', r'\1 \2', text)
    # Remaining script boundaries after look-alike letters were fixed.
    text = re.sub(r'(?<=[А-Яа-яЁё])(?=[A-Za-z])|(?<=[A-Za-z])(?=[А-Яа-яЁё])', ' ', text)
    text = re.sub(r'(?<=\S)\(', ' (', text)
    text = re.sub(r'\)(?=\S)', ') ', text)
    return text


def year(value):
    number = int(value)
    if number < 100:
        number += 1900 if number >= 80 else 2000
    return number if 1980 <= number <= MAX_YEAR else None


def classify(name):
    """Return a dict with every attribute that can be read from one product name."""
    text = normalize(name)
    location = []
    shelf = SHELF.match(text)
    if shelf:
        location.append(f'{shelf.group(1).upper()}{shelf.group(2)}-{shelf.group(3)}')
        text = text[shelf.end():].strip()
    for box in BOX.finditer(text):
        location.append(f'Ящик {box.group(1)}')
    text = BOX.sub(' ', text).strip()
    text = re.sub(r'\s+', ' ', text)
    lower = text.lower()

    group, part = None, None
    for pattern, code, label in PART_RULES:
        if re.search(pattern, lower):
            group, part = code, label
            break

    raw = TOKEN.findall(text)
    tokens = []
    i = 0
    while i < len(raw):
        pair = (raw[i].upper(), raw[i + 1].upper()) if i + 1 < len(raw) else None
        if pair in COMPOUNDS:
            tokens.append(COMPOUNDS[pair])
            i += 2
            continue
        tokens.append(raw[i])
        i += 1

    fitments = []
    makes, brands, countries, unknown, engines = [], [], [], [], []
    loose_years = []
    in_parens = False
    last_kind = None
    group_start = 0  # fitments listed since the previous year share the next year range

    def new_fitment(make, model=None, generation=None):
        fitments.append({'make': make, 'model': model, 'generation': generation,
                         'year_from': None, 'year_to': None, 'engine': None})
        return fitments[-1]

    for index, token in enumerate(tokens):
        upper = token.upper().strip("'.,")
        if token == '(':
            in_parens = True
            continue
        if token == ')':
            in_parens = False
            continue
        if not re.search(r'[A-Za-z0-9]', token):
            continue
        # Years: «11-», «98-05», «2020-», «(-2011)».
        range_match = YEAR_RANGE.match(upper)
        to_match = TO_YEAR.match(upper)
        if range_match and year(range_match.group(1)):
            start = year(range_match.group(1))
            end = year(range_match.group(2)) if range_match.group(2) else None
            if end is not None and end < start:
                end = None
            targets = [f for f in fitments[group_start:] if f['year_from'] is None and f['year_to'] is None]
            for target in targets:
                target['year_from'], target['year_to'] = start, end
            if not targets:
                loose_years.append((start, end))
            group_start = len(fitments)
            last_kind = 'year'
            continue
        if to_match and year(to_match.group(1)):
            targets = [f for f in fitments[group_start:] if f['year_to'] is None] or (
                [fitments[-1]] if fitments and fitments[-1]['year_to'] is None else [])
            for target in targets:
                target['year_to'] = year(to_match.group(1))
            if not targets:
                loose_years.append((None, year(to_match.group(1))))
            group_start = len(fitments)
            last_kind = 'year'
            continue
        if ENGINE_CODE.match(upper) or upper in ENGINE_FAMILIES:
            engines.append(upper)
            if fitments and fitments[-1]['engine'] is None:
                fitments[-1]['engine'] = upper
            last_kind = 'engine'
            continue
        if VOLUME.match(upper):
            if last_kind in {'model', 'generation', 'engine', 'year'}:
                engines.append(upper)
            continue
        if upper in MODELS:
            make, model = MODELS[upper]
            new_fitment(make, model)
            makes.append(make)
            last_kind = 'model'
            continue
        if upper in GENERATIONS and (last_kind in {'model', 'generation'} or in_parens
                                     or upper in DISTINCT_GENERATIONS):
            make, model = GENERATIONS[upper]
            last = fitments[-1] if fitments else None
            if last and last['generation'] is None and (last['model'] == model or last_kind == 'model'):
                last['generation'] = upper
            else:
                new_fitment(make, model, upper)
            makes.append(make)
            last_kind = 'generation'
            continue
        if upper in MAKE_TOKENS:
            makes.append(MAKE_TOKENS[upper])
            last_kind = 'make'
            continue
        if upper in COUNTRIES:
            countries.append(COUNTRIES[upper])
            continue
        if not re.search(r'[A-Za-z]', upper):
            continue
        parts = [p for p in re.split(r'[-/]', upper) if p]
        if upper in BRANDS:
            if BRANDS[upper]:
                brands.append(BRANDS[upper])
            last_kind = 'brand'
            continue
        # «LEON-HPT», «KOREA-TENACITY», «SPEEDMATE-JD»: known brands plus short unknown suffixes.
        if parts and any(p in BRANDS and BRANDS[p] for p in parts) and all(
                p in BRANDS or p in MAKE_TOKENS or p in COUNTRIES or p.isdigit() or len(p) <= 3 for p in parts):
            for p in parts:
                if p in BRANDS and BRANDS[p]:
                    brands.append(BRANDS[p])
                elif p in MAKE_TOKENS:
                    makes.append(MAKE_TOKENS[p])
                elif p in COUNTRIES:
                    countries.append(COUNTRIES[p])
            last_kind = 'brand'
            continue
        if upper in NOISE or upper in GENERATIONS or len(upper) < 2 or re.search(r'\d', upper):
            continue
        unknown.append(upper)

    makes = list(dict.fromkeys(makes))
    if not fitments and makes:
        for make in makes:
            new_fitment(make)
    if loose_years and fitments:
        for fitment in fitments:
            if fitment['year_from'] is None and fitment['year_to'] is None:
                fitment['year_from'], fitment['year_to'] = loose_years[0]
    years_from = [f['year_from'] for f in fitments if f['year_from']] + [y[0] for y in loose_years if y[0]]
    years_to = [f['year_to'] for f in fitments if f['year_to']] + [y[1] for y in loose_years if y[1]]

    position = []
    if re.search(r'передн|\bперед\b|\bпер\b|\bпер\.|^пер', lower):
        position.append('передний')
    if re.search(r'задн|\bзад\b', lower):
        position.append('задний')
    if re.search(r'\bлев|\bl\b|\blh\b', lower):
        position.append('левый')
    if re.search(r'\bправ|\br\b|\brh\b', lower):
        position.append('правый')
    if re.search(r'верхн', lower):
        position.append('верхний')
    if re.search(r'нижн', lower):
        position.append('нижний')

    # The part phrase: Cyrillic words before the first Latin brand/vehicle token.
    phrase = re.split(r'\s(?=[A-Za-z])', text, maxsplit=1)[0]
    phrase = re.sub(r'\s+[\d*.,xх+]+(?:\s?мм|mm)?$', '', phrase).strip(' ,.-')

    issues = []
    if not group:
        issues.append('категория')
    if not fitments:
        issues.append('марка/модель')
    elif not any(f['model'] for f in fitments):
        issues.append('модель')
    if fitments and not years_from and not years_to:
        issues.append('годы')
    return {
        'name': str(name).strip(),
        'part_name': phrase[:1].upper() + phrase[1:] if phrase else '',
        'location': ', '.join(location),
        'group': group,
        'category': GROUPS[group.split('-')[0]] if group else None,
        'subcategory': GROUPS[group] if group else None,
        'part': part,
        'position': ', '.join(position),
        'brands': list(dict.fromkeys(brands)),
        'countries': list(dict.fromkeys(countries)),
        'makes': list(dict.fromkeys(f['make'] for f in fitments)) if fitments else makes,
        'fitments': fitments,
        'year_from': min(years_from) if years_from else None,
        'year_to': max(years_to) if years_to else None,
        'engines': list(dict.fromkeys(engines)),
        'unknown': unknown,
        'issues': issues,
    }


def fitment_text(f):
    model = ' '.join(x for x in [f['make'], f['model'], f'({f["generation"]})' if f['generation'] else None] if x)
    if f['year_from'] or f['year_to']:
        model += f" {f['year_from'] or ''}–{f['year_to'] or ''}".replace('– ', '–')
    return model


def main(source, target):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = load_workbook(source, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    col = {name: header.index(name) for name in header}
    products = []
    for row in rows:
        if not row or not row[col['ID']]:
            continue
        result = classify(row[col['Наименование']])
        result.update(id=row[col['ID']], code=row[col['Код']], sku=row[col['Артикул']],
                      unit=row[col['Единица измерения']],
                      price=row[col.get('СКЛАД (Цена)')], stock=row[col.get('СКЛАД (Остаток)')])
        products.append(result)

    book = Workbook()
    font = Font(name='Arial', size=10)
    bold = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1F3A5F')
    warn = PatternFill('solid', fgColor='FFF4CC')

    def write(ws, headers, data, widths, flag=None):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        for values in data:
            ws.append(values)
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.font = font
        if flag is not None:
            for index, values in enumerate(data, start=2):
                if values[flag]:
                    for cell in ws[index]:
                        cell.fill = warn
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    ws = book.active
    ws.title = 'Товары'
    write(ws, ['ID CloudShop', 'Код', 'Артикул', 'Наименование (исходное)', 'Деталь', 'Ячейка склада',
               'Код группы', 'Категория', 'Подкатегория', 'Тип детали', 'Положение', 'Бренд', 'Страна',
               'Марка', 'Применимость', 'Год от', 'Год до', 'Двигатель', 'Цена СКЛАД, ₸', 'Остаток СКЛАД',
               'Проверить'],
          [[p['id'], p['code'], p['sku'], p['name'], p['part_name'], p['location'], p['group'], p['category'],
            p['subcategory'], p['part'], p['position'], ', '.join(p['brands']), ', '.join(p['countries']),
            ', '.join(p['makes']), '; '.join(fitment_text(f) for f in p['fitments']), p['year_from'],
            p['year_to'], ', '.join(p['engines']), p['price'], p['stock'], ', '.join(p['issues'])]
           for p in products],
          [26, 9, 22, 60, 30, 12, 9, 22, 28, 26, 18, 20, 10, 16, 50, 8, 8, 14, 12, 10, 22], flag=20)

    write(book.create_sheet('Применимость'),
          ['ID CloudShop', 'Наименование', 'Марка', 'Модель', 'Кузов', 'Год от', 'Год до', 'Двигатель',
           'Положение', 'Источник', 'Статус'],
          [[p['id'], p['name'], f['make'], f['model'], f['generation'], f['year_from'], f['year_to'],
            f['engine'], p['position'], 'название товара', 'авто, проверить']
           for p in products for f in p['fitments']],
          [26, 60, 12, 16, 8, 8, 8, 10, 18, 16, 16])

    counts = Counter(p['group'] for p in products if p['group'])
    write(book.create_sheet('Категории'), ['Код группы', 'Категория', 'Подкатегория', 'Товаров'],
          [[code, GROUPS[code.split('-')[0]], GROUPS[code] if '-' in code else '',
            sum(v for k, v in counts.items() if k.startswith(code)) if '-' not in code else counts.get(code, 0)]
           for code in GROUPS],
          [12, 28, 34, 10])

    brand_counts = Counter(b for p in products for b in p['brands'])
    write(book.create_sheet('Бренды'), ['Бренд', 'Товаров'], brand_counts.most_common(), [24, 10])

    model_counts = Counter((f['make'], f['model'] or '(без модели)') for p in products for f in p['fitments'])
    write(book.create_sheet('Модели'), ['Марка', 'Модель', 'Товаров'],
          [[m, mo, n] for (m, mo), n in sorted(model_counts.items(), key=lambda x: (x[0][0], -x[1]))],
          [14, 18, 10])

    unknown_counts = Counter(t for p in products for t in p['unknown'])
    examples = {}
    for p in products:
        for t in p['unknown']:
            examples.setdefault(t, p['name'])
    write(book.create_sheet('Нераспознанные слова'), ['Слово', 'Встречается', 'Пример наименования'],
          [[t, n, examples[t]] for t, n in unknown_counts.most_common()], [18, 12, 70])

    summary = book.create_sheet('Сводка', 0)
    total = len(products)
    stats = [
        ('Товаров в выгрузке', total),
        ('С категорией', sum(1 for p in products if p['group'])),
        ('С маркой авто', sum(1 for p in products if p['fitments'])),
        ('С моделью авто', sum(1 for p in products if any(f['model'] for f in p['fitments']))),
        ('С годами выпуска', sum(1 for p in products if p['year_from'] or p['year_to'])),
        ('С брендом детали', sum(1 for p in products if p['brands'])),
        ('С ячейкой склада', sum(1 for p in products if p['location'])),
        ('Без замечаний', sum(1 for p in products if not p['issues'])),
    ]
    summary.append(['Показатель', 'Товаров', 'Доля'])
    for label, value in stats:
        summary.append([label, value, value / total if total else 0])
    summary.append([])
    for line in [
        'Все поля получены из наименования товара: в выгрузке CloudShop категория, бренд и страна не заполнены.',
        'Это предложение для проверки. Строки, выделенные жёлтым на листе «Товары», требуют ручной проверки (колонка «Проверить»).',
        'Дерево категорий — docs/cloudshop/templates/catalog_groups.csv, дополнено группами G01-05, G05-05, G05-06, G08-04, G08-05.',
        'Применимость: одна строка на автомобиль; годы «11-» означают «с 2011», «98-05» — «1998–2005».',
        'Себестоимость и закупочные цены из выгрузки не переносились.',
        'Пересчёт: python3 tools/catalog/classify_products.py <выгрузка.xlsx> <результат.xlsx>',
    ]:
        summary.append([line])
    for r in summary.iter_rows():
        for cell in r:
            cell.font = font
    for cell in summary[1]:
        cell.font, cell.fill = bold, fill
    for r in summary.iter_rows(min_row=2, max_row=len(stats) + 1):
        r[2].number_format = '0.0%'
    summary.column_dimensions['A'].width = 30
    summary.column_dimensions['B'].width = 12
    summary.column_dimensions['C'].width = 10
    book.save(target)
    for label, value in stats:
        print(f'{label}: {value} ({value / total:.1%})')
    print('Нераспознанные слова (топ):', ', '.join(f'{t}×{n}' for t, n in unknown_counts.most_common(40)))
    print('Результат:', target)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
